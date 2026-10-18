"""PDF conversion engine supporting conversion to DOCX, PPTX, XLSX, JPG, PNG, EPUB, TXT, MD."""

from __future__ import annotations

import logging
import re
import shutil
from pathlib import Path
from typing import Any, Dict

import pandas as pd
from pdf2docx import Converter as Pdf2DocxConverter
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from filetool.core.deps import resolver
from filetool.core.registry import Converter, OptionDefinition, registry
from filetool.core.utils import run_command, safe_temp_dir

logger = logging.getLogger("filetool.engines.pdf")

PDF_OPTIONS = {
    "dpi": OptionDefinition(
        name="dpi",
        label="Render DPI",
        type=int,
        default=300,
        min_value=72,
        max_value=600,
        description="Resolution for rendering PDF pages to images or slides",
    ),
}


def convert_pdf_to_docx(src: Path, dest: Path, options: Dict[str, Any]) -> Path:
    """Convert PDF to DOCX using pdf2docx."""
    cv = Pdf2DocxConverter(str(src))
    try:
        cv.convert(str(dest))
    finally:
        cv.close()

    if not dest.exists():
        raise RuntimeError(f"PDF to DOCX conversion failed to produce {dest}")
    return dest


def convert_pdf_to_pptx(src: Path, dest: Path, options: Dict[str, Any]) -> Path:
    """
    Convert PDF to PPTX presentation.
    Renders each PDF page at high DPI using pdftoppm, then creates a PowerPoint slide for each page.
    """
    pdftoppm = resolver.get_tool_path("pdftoppm")
    dpi = int(options.get("dpi", 300))

    with safe_temp_dir(prefix="filetool_pdf_pptx_") as tmp_dir:
        prefix = tmp_dir / "page"
        render_cmd = [pdftoppm, "-png", "-r", str(dpi), str(src), str(prefix)]
        res = run_command(render_cmd)
        if not res.success:
            raise RuntimeError(f"Rendering PDF pages for PPTX failed: {res.stderr_tail()}")

        page_images = sorted(tmp_dir.glob("page-*.png"))
        if not page_images:
            raise RuntimeError("No rendered page images found for PPTX generation")

        prs = Presentation()
        # Use dimensions from the first page image
        with Image.open(page_images[0]) as first_img:
            width_px, height_px = first_img.size
            aspect = width_px / height_px

        # Standard slide sizing
        if aspect > 1.2:
            prs.slide_width = Inches(10)
            prs.slide_height = Inches(7.5)
        else:
            prs.slide_width = Inches(8.5)
            prs.slide_height = Inches(11)

        blank_layout = prs.slide_layouts[6]
        for img_path in page_images:
            slide = prs.slides.add_slide(blank_layout)
            slide.shapes.add_picture(
                str(img_path),
                0,
                0,
                width=prs.slide_width,
                height=prs.slide_height,
            )

        prs.save(str(dest))

    return dest


def convert_pdf_to_xlsx(src: Path, dest: Path, options: Dict[str, Any]) -> Path:
    """
    Convert PDF to XLSX spreadsheet (best effort extraction).
    Extracts tabular layout text with pdftotext, parses delimited blocks into columns, and outputs Excel.
    """
    pdftotext = resolver.get_tool_path("pdftotext")

    with safe_temp_dir(prefix="filetool_pdf_xlsx_") as tmp_dir:
        txt_out = tmp_dir / "layout.txt"
        cmd = [pdftotext, "-layout", str(src), str(txt_out)]
        res = run_command(cmd)
        if not res.success or not txt_out.exists():
            raise RuntimeError(f"pdftotext extraction failed: {res.stderr_tail()}")

        lines = txt_out.read_text(encoding="utf-8", errors="replace").splitlines()
        rows = []
        for line in lines:
            line_str = line.strip()
            if not line_str:
                continue
            # Split by 2 or more spaces to detect table columns
            cols = [col.strip() for col in re.split(r"\s{2,}", line_str) if col.strip()]
            if cols:
                rows.append(cols)

        if not rows:
            rows = [["No text or tables detected in PDF"]]

        # Standardize column count
        max_cols = max(len(r) for r in rows)
        padded_rows = [r + [""] * (max_cols - len(r)) for r in rows]

        df = pd.DataFrame(padded_rows)
        df.to_excel(dest, index=False, header=False, engine="openpyxl")

    return dest


def convert_pdf_to_images(src: Path, dest: Path, options: Dict[str, Any]) -> Path:
    """
    Convert PDF to JPG or PNG images (300 DPI, all pages).
    First page is saved to dest; subsequent pages are saved alongside dest as 'name (Page 2).ext'.
    """
    pdftoppm = resolver.get_tool_path("pdftoppm")
    out_ext = dest.suffix.lower().lstrip(".")
    dpi = int(options.get("dpi", 300))

    with safe_temp_dir(prefix="filetool_pdf_img_") as tmp_dir:
        fmt_flag = "-png" if out_ext == "png" else "-jpeg"
        prefix = tmp_dir / "page"
        cmd = [pdftoppm, fmt_flag, "-r", str(dpi), str(src), str(prefix)]
        res = run_command(cmd)
        if not res.success:
            raise RuntimeError(f"pdftoppm rendering failed: {res.stderr_tail()}")

        rendered_pages = sorted(tmp_dir.glob("page-*.*"))
        if not rendered_pages:
            raise RuntimeError("No rendered page images produced from PDF")

        shutil.copy2(rendered_pages[0], dest)

        if len(rendered_pages) > 1:
            dest_dir = dest.parent
            dest_stem = dest.stem
            for idx, img_path in enumerate(rendered_pages[1:], start=2):
                sibling_dest = dest_dir / f"{dest_stem} (Page {idx}).{out_ext}"
                shutil.copy2(img_path, sibling_dest)

    return dest


def convert_pdf_to_txt(src: Path, dest: Path, options: Dict[str, Any]) -> Path:
    """Convert PDF to text preserving layout using pdftotext."""
    pdftotext = resolver.get_tool_path("pdftotext")
    cmd = [pdftotext, "-layout", str(src), str(dest)]
    res = run_command(cmd)
    if not res.success or not dest.exists():
        raise RuntimeError(f"PDF text extraction failed: {res.stderr_tail()}")
    return dest


def convert_pdf_to_md(src: Path, dest: Path, options: Dict[str, Any]) -> Path:
    """Convert PDF to clean Markdown format."""
    pdftotext = resolver.get_tool_path("pdftotext")

    with safe_temp_dir(prefix="filetool_pdf_md_") as tmp_dir:
        txt_path = tmp_dir / "extracted.txt"
        cmd = [pdftotext, "-layout", str(src), str(txt_path)]
        res = run_command(cmd)
        if not res.success or not txt_path.exists():
            raise RuntimeError(f"PDF extraction failed: {res.stderr_tail()}")

        content = txt_path.read_text(encoding="utf-8", errors="replace")
        # Format pages as markdown sections
        pages = content.split("\x0c")  # form-feed page separator
        md_lines = []
        for i, page in enumerate(pages, start=1):
            stripped = page.strip()
            if not stripped:
                continue
            if len(pages) > 1:
                md_lines.append(f"## Page {i}\n")
            md_lines.append(stripped)
            md_lines.append("\n")

        dest.write_text("\n".join(md_lines), encoding="utf-8")

    return dest


def convert_pdf_to_epub(src: Path, dest: Path, options: Dict[str, Any]) -> Path:
    """Convert PDF to EPUB using pandoc on intermediate extracted text."""
    pandoc = resolver.get_tool_path("pandoc")

    with safe_temp_dir(prefix="filetool_pdf_epub_") as tmp_dir:
        md_path = tmp_dir / "intermediate.md"
        convert_pdf_to_md(src, md_path, options)

        cmd = [pandoc, str(md_path), "-f", "markdown", "-t", "epub", "-o", str(dest)]
        res = run_command(cmd)
        if not res.success or not dest.exists():
            raise RuntimeError(f"Pandoc EPUB generation from PDF failed: {res.stderr_tail()}")

    return dest


def register_pdf_converters() -> None:
    """Register all PDF conversions."""
    # PDF to DOCX
    registry.register(
        Converter(
            id="pdf_to_docx",
            name="PDF to DOCX",
            input_formats=("pdf",),
            output_format="docx",
            category="PDF",
            convert_fn=convert_pdf_to_docx,
            description="Convert PDF document to editable Microsoft Word DOCX",
        )
    )

    # PDF to PPTX
    registry.register(
        Converter(
            id="pdf_to_pptx",
            name="PDF to PPTX",
            input_formats=("pdf",),
            output_format="pptx",
            category="PDF",
            convert_fn=convert_pdf_to_pptx,
            options=PDF_OPTIONS,
            required_tools=("pdftoppm",),
            description="Convert PDF to PowerPoint presentation (one slide per page)",
        )
    )

    # PDF to XLSX
    registry.register(
        Converter(
            id="pdf_to_xlsx",
            name="PDF to XLSX",
            input_formats=("pdf",),
            output_format="xlsx",
            category="PDF",
            convert_fn=convert_pdf_to_xlsx,
            required_tools=("pdftotext",),
            description="Extract tabular data from PDF to Microsoft Excel XLSX (best effort)",
        )
    )

    # PDF to PNG
    registry.register(
        Converter(
            id="pdf_to_png",
            name="PDF to PNG",
            input_formats=("pdf",),
            output_format="png",
            category="PDF",
            convert_fn=convert_pdf_to_images,
            options=PDF_OPTIONS,
            required_tools=("pdftoppm",),
            description="Render PDF pages to 300 DPI PNG images",
        )
    )

    # PDF to JPG
    registry.register(
        Converter(
            id="pdf_to_jpg",
            name="PDF to JPG",
            input_formats=("pdf",),
            output_format="jpg",
            category="PDF",
            convert_fn=convert_pdf_to_images,
            options=PDF_OPTIONS,
            required_tools=("pdftoppm",),
            description="Render PDF pages to 300 DPI JPG images",
        )
    )

    # PDF to TXT
    registry.register(
        Converter(
            id="pdf_to_txt",
            name="PDF to TXT",
            input_formats=("pdf",),
            output_format="txt",
            category="PDF",
            convert_fn=convert_pdf_to_txt,
            required_tools=("pdftotext",),
            description="Extract plain text from PDF preserving visual layout",
        )
    )

    # PDF to MD
    registry.register(
        Converter(
            id="pdf_to_md",
            name="PDF to Markdown",
            input_formats=("pdf",),
            output_format="md",
            category="PDF",
            convert_fn=convert_pdf_to_md,
            required_tools=("pdftotext",),
            description="Extract structured text from PDF into Markdown",
        )
    )

    # PDF to EPUB
    registry.register(
        Converter(
            id="pdf_to_epub",
            name="PDF to EPUB",
            input_formats=("pdf",),
            output_format="epub",
            category="PDF",
            convert_fn=convert_pdf_to_epub,
            required_tools=("pdftotext", "pandoc"),
            description="Convert PDF content into EPUB eBook format",
        )
    )


# Automatically register on import
register_pdf_converters()
